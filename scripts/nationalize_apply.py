"""历史 14 列版本全国化改造（Excel + md 同步，硬约束：不增删行、不调序）

仅用于归档版本复现，不操作当前前台的 15 列源文件。

改造规则（对应扫描报告的 31 条 TEXT_FIX + 3,020 条 REGION_FIX）：
1. 「适用区域」列全部非"全国"值 → "全国"（含细分省份值）
2. 补贴/补助参考金额、省名限定 → 通用表述（"以当地当年官方公示为准"）
3. 指标名称中的地区词 → 泛化或去除（"东北民俗工艺品产量"→"区域民俗工艺品产量"）
4. 风险含义/取值说明中的省名、长白山、北大荒等地域词 → 去除或泛化
5. 特色语义词（黑土地/寒地/人参/冷水鱼/亚麻/林下参等）**保留**（地域特色模块）

产物（不覆盖原文件）：
- docs_archive_20260920/农业及相关产业动态指标搜集体系_national.xlsx
- docs_archive_20260920/农业及相关产业动态指标搜集体系_national.md
- back-end/data/nationalize_changes.csv（逐单元格变更明细）

用法：
    python scripts/nationalize_apply.py
"""

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs_archive_20260920"
XLSX = DOCS / "农业及相关产业动态指标搜集体系.xlsx"
MD = DOCS / "农业及相关产业动态指标搜集体系.md"
XLSX_OUT = DOCS / "农业及相关产业动态指标搜集体系_national.xlsx"
MD_OUT = DOCS / "农业及相关产业动态指标搜集体系_national.md"
CHANGES = ROOT / "data" / "nationalize_changes.csv"

# ---------- 1. 适用区域：所有非"全国"值 → 全国 ----------
REGION_OLD = [
    "东北全境",
    "黑吉辽+蒙东",
    "黑龙江+吉林",
    "黑龙江",
    "黑龙江+蒙东",
    "辽宁",
    "吉林",
    "黑龙江+内蒙古+吉林（农垦/森工体系）",
    "仅辽宁鞍山",
    "黑龙江+辽宁",
    "辽宁+吉林",
    "吉林+辽宁",
    "仅辽宁盘锦",
]
REGION_MAP = {v: "全国" for v in REGION_OLD}

# ---------- 2. 单元格文本改写（整格精确匹配，原文 → 新文） ----------
CELL_MAP = {
    # 农垦体系主体（去"北大荒/呼伦贝尔"示例）
    "农垦（北大荒等）体系内主体": "农垦/森工体系内主体",
    "是否属农垦集团体系（北大荒/呼伦贝尔农垦等）": "是否属农垦/森工集团体系",
    # 补贴/补助参考金额 → 通用表述
    (
        "年度实际到账金额（东北近年参考：黑75.73、吉约128、辽约132元/亩，以当年官方公示为准）"
    ): "年度实际到账金额（以当地当年官方公示为准）",
    (
        "免耕/少耕/秸秆覆盖还田实施面积（黑龙江补助约35-60元/亩、吉林最高约80元/亩，以当年官方公示为准）"
    ): "免耕/少耕/秸秆覆盖还田实施面积（以当地当年官方公示为准）",
    "深松作业面积（黑龙江作业补助约20元/亩，以当年官方公示为准）": "深松作业面积（以当地当年官方公示为准）",
    (
        "当年补贴标准×种植面积（东北近年参考：黑约30-150、吉约49、辽约67元/亩，黑龙江当年标准以省级公示为准）"
    ): "当年补贴标准×种植面积（以当地当年官方公示为准）",
    (
        "当年补贴标准×种植面积（东北近年参考：黑地表水约157、地下水约107元/亩，以当年官方公示为准）"
    ): "当年补贴标准×种植面积（以当地当年官方公示为准）",
    (
        "轮作/休耕试点补贴到账（东北近年参考：轮作约150-200元/亩、休耕约500-800元/亩·季，以当年官方公示为准）"
    ): "轮作/休耕试点补贴到账（以当地当年官方公示为准）",
    (
        "当年补贴标准×面积（东北近年参考：黑约350、吉约550、辽约410-460元/亩；"
        "高油/高蛋白品种另有额外补贴，以当年官方公示为准）"
    ): "当年补贴标准×面积（高油/高蛋白品种另有额外补贴，以当地当年官方公示为准）",
    "当年平均含糖率（东北近年参考约16-18%）": "当年平均含糖率（参考区间约16-18%，以当地实际为准）",
    # 风险含义 / 取值说明：去省名等地名（特色词保留）
    "黑龙江亚麻原料保障": "亚麻原料保障",
    "长白山人参产业规模与周期": "人参产业规模与周期",
    "长白山人参价值": "人参价值",
    "黑龙江名优渔产": "名优渔产",
    "东北产区原料采购占比": "本地及区域原料采购占比",
    "玉米深加工规模（东北支柱）": "玉米深加工规模（区域支柱）",
    "高端乳品（黑龙江）": "高端乳品",
    "东北杂粮特色": "杂粮特色",
    "东北杂粮加工": "杂粮加工",
    "黑龙江大豆发酵特色": "寒地大豆发酵特色",
    "长白山药材加工": "药材深加工",
    "黑龙江亚麻特色": "亚麻加工特色",
    "原料保障（黑龙江亚麻）": "亚麻原料保障",
    "东北林区林业": "寒地林区林业",
    "东北民俗特色": "民俗特色",
    "东北林化产业": "林化产业",
    "东北大农机产业": "大农机产业",
    "东北粮贸规模": "粮食贸易规模",
    # 指标名称：去地区词（泛化）
    "东北民俗工艺品产量": "区域民俗工艺品产量",
    "东北农业务工劳务服务规模": "农业务工劳务服务规模",
    "东北农村人居环境整治覆盖村数": "农村人居环境整治覆盖村数",
    "东北农产品品牌广告服务额": "区域农产品品牌广告服务额",
}

# ---------- 3. md 头部说明改写（整行替换） ----------
MD_HEAD_REPL = [
    ("# 农业及相关产业分层动态指标搜集体系（东北振兴版）", "# 农业及相关产业分层动态指标搜集体系（全国版）"),
    (
        "> 依据《农业及相关产业统计分类（2020）》编制 · 面向**东北振兴区域**涉农信贷风险评估 · 树形分层",
        "> 依据《农业及相关产业统计分类（2020）》编制 · 面向**全国范围**涉农信贷风险评估 · 树形分层",
    ),
    (
        "**适用范围**：面向**东北振兴区域**（黑龙江、吉林、辽宁及内蒙古东部）内的"
        "**涉农小微企业、农民专业合作社和个体工商户**（不含农户个人信贷）。"
        "指标按**树形层级**组织（基本项 → 大类 → 中类 → 小类）。",
        "**适用范围**：面向**全国范围**内的**涉农小微企业、农民专业合作社和个体工商户**"
        "（不含农户个人信贷）。指标按**树形层级**组织（基本项 → 大类 → 中类 → 小类）。",
    ),
    (
        '**适用区域**：「适用区域」列标注东北区域内细分范围（"东北全境"或"黑龙江""吉林""辽宁""黑龙江+吉林""黑吉辽+蒙东"等；蒙东五盟市属东北振兴范围，与黑吉辽统一口径）。',
        '**适用区域**：「适用区域」列统一为"全国"；原区域细分口径已并入通用表述，各地可按当地政策落地实施。',
    ),
    (
        "**特色指标**：「是否特色指标」= 是的为**东北/寒地区域特色**指标"
        '（黑土地、寒地、人参、冷水鱼、冰雪等），通用经营指标标"否"。',
        "**特色指标**：「是否特色指标」= 是的为**地域特色**指标"
        "（黑土地、寒地、人参、冷水鱼、冰雪、亚麻等地域特征，作为地域特色模块供各地按需选用），"
        '通用经营指标标"否"。',
    ),
]


def apply_text(v: str) -> tuple[str, str]:
    """返回 (新值, 规则标记)"""
    if v in REGION_MAP:
        return REGION_MAP[v], "REGION"
    if v in CELL_MAP:
        return CELL_MAP[v], "TEXT"
    return v, ""


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    changes: list[dict] = []

    # ---------- Excel ----------
    wb = load_workbook(XLSX)  # 保留样式
    print(f"Excel sheets: {wb.sheetnames}")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or not isinstance(cell.value, str):
                    continue
                nv, rule = apply_text(cell.value)
                if rule:
                    changes.append(
                        {
                            "file": "xlsx",
                            "sheet": ws.title,
                            "cell": cell.coordinate,
                            "rule": rule,
                            "old": cell.value,
                            "new": nv,
                        }
                    )
                    cell.value = nv
    wb.save(XLSX_OUT)
    print(f"已写出: {XLSX_OUT.name}")

    # ---------- md ----------
    # 注意：md 必须按「表格单元格」整格匹配，不能用全局字符串替换
    # （否则"黑龙江+蒙东"会被"黑龙江"先替换成"全国+蒙东"，or 长文本内误伤）
    text = MD.read_text(encoding="utf-8")
    lines = text.split("\n")
    head_changed = 0
    md_cells = 0
    out_lines: list[str] = []
    for line in lines:
        stripped = line.strip()
        # 头部说明整行替换
        head_hit = False
        for old, new in MD_HEAD_REPL:
            if stripped == old:
                out_lines.append(new)
                head_changed += 1
                head_hit = True
                break
        if head_hit:
            continue
        # 表格行：按 | 分格整格匹配
        if stripped.startswith("|") and stripped.count("|") >= 3:
            parts = line.split("|")
            for i, seg in enumerate(parts):
                s = seg.strip()
                if s in REGION_MAP:
                    parts[i] = seg.replace(s, REGION_MAP[s])
                    md_cells += 1
                elif s in CELL_MAP:
                    parts[i] = seg.replace(s, CELL_MAP[s])
                    md_cells += 1
            out_lines.append("|".join(parts))
        else:
            out_lines.append(line)
    text = "\n".join(out_lines)
    MD_OUT.write_text(text, encoding="utf-8")
    print(f"已写出: {MD_OUT.name}（头部替换 {head_changed}/{len(MD_HEAD_REPL)}，单元格替换 {md_cells} 处）")
    changes.append(
        {
            "file": "md",
            "sheet": "-",
            "cell": "summary",
            "rule": "MD",
            "old": f"head={head_changed}",
            "new": f"cells={md_cells}",
        }
    )

    # ---------- 变更明细 ----------
    CHANGES.parent.mkdir(parents=True, exist_ok=True)
    with CHANGES.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["file", "sheet", "cell", "rule", "old", "new"])
        w.writeheader()
        w.writerows(changes)
    print(f"变更明细: {CHANGES}（共 {len(changes)} 条）")


if __name__ == "__main__":
    main()
